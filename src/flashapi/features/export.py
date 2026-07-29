"""Export functionality — CSV (built-in), XLSX and PDF (optional deps)."""

from __future__ import annotations

import csv
import io
from typing import Any


def export_csv(items: list[dict[str, Any]], fields: list[str]) -> bytes:
    output = io.BytesIO()
    output.write(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel compatibility
    wrapper = io.TextIOWrapper(output, encoding="utf-8", newline="")
    writer = csv.DictWriter(wrapper, fieldnames=fields, extrasaction="ignore", delimiter=";")
    writer.writeheader()
    for item in items:
        writer.writerow({k: item.get(k, "") for k in fields})
    wrapper.flush()
    wrapper.detach()
    return output.getvalue()


def export_xlsx(items: list[dict[str, Any]], fields: list[str]) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl is required for XLSX export: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, field in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, item in enumerate(items, 2):
        for col_idx, field in enumerate(fields, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=item.get(field, ""))
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col_idx, field in enumerate(fields, 1):
        max_length = len(str(field))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 3, 50)

    ws.auto_filter.ref = ws.dimensions

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_pdf(items: list[dict[str, Any]], fields: list[str]) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
    except ImportError:
        raise ImportError("reportlab is required for PDF export: pip install reportlab")

    output = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    title = Paragraph("Export", styles["Title"])
    elements.append(title)
    elements.append(Spacer(1, 5 * mm))

    table_data = [fields]
    for item in items:
        row = [str(item.get(f, "")) if item.get(f) is not None else "" for f in fields]
        table_data.append(row)

    available_width = page_size[0] - 20 * mm
    col_count = len(fields)
    col_width = available_width / col_count if col_count else available_width

    table = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))

    elements.append(table)
    doc.build(elements)
    return output.getvalue()


EXPORTERS = {
    "csv": export_csv,
    "xlsx": export_xlsx,
    "pdf": export_pdf,
}

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
